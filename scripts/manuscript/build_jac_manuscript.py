from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
ASSET_DIR = ROOT / "manuscript_assets" / "phase_finder"
OUTPUT = Path("/private/tmp/Phase finder_JAC_revised.docx")

BLUE = RGBColor(31, 78, 121)
GREEN = RGBColor(38, 110, 82)
RED = RGBColor(175, 55, 55)
GREY = RGBColor(90, 100, 110)
LIGHT_BLUE = "EAF2F8"
LIGHT_YELLOW = "FFF2CC"
LIGHT_GREY = "F2F4F6"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def set_paragraph_border(paragraph, color: str = "D9E2F3", size: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_run(paragraph, text: str, *, bold=False, italic=False, color=None, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if size is not None:
        run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    return run


def add_hyperlink(paragraph, text: str, url: str, *, size=9):
    relationship_id = paragraph.part.relate_to(
        url,
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    font = OxmlElement("w:rFonts")
    font.set(qn("w:ascii"), "Times New Roman")
    font.set(qn("w:hAnsi"), "Times New Roman")
    properties.append(font)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), str(int(size * 2)))
    properties.append(font_size)
    run.append(properties)

    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_text_with_links(paragraph, text: str, *, size=9):
    cursor = 0
    for match in re.finditer(r"https?://[^\s]+", text):
        start, end = match.span()
        raw_url = match.group(0)
        url = raw_url.rstrip(".,;)")
        suffix = raw_url[len(url):]
        if start > cursor:
            add_run(paragraph, text[cursor:start], size=size)
        add_hyperlink(paragraph, url, url, size=size)
        if suffix:
            add_run(paragraph, suffix, size=size)
        cursor = end
    if cursor < len(text):
        add_run(paragraph, text[cursor:], size=size)


def add_body(doc: Document, text: str, *, italic=False, keep=False):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    if "http://" in text or "https://" in text:
        add_text_with_links(p, text, size=10.5)
    else:
        add_run(p, text, italic=italic)
    if keep:
        keep_with_next(p)
    return p


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(4)
    run = add_run(p, text, bold=True, color=BLUE, size=13 if level == 1 else 11.5)
    return p


def add_equation(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    add_run(p, text, italic=True, size=10.5)
    return p


def add_author_action(doc: Document, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(15.8)
    cell = table.cell(0, 0)
    cell.width = Cm(15.8)
    set_cell_shading(cell, LIGHT_YELLOW)
    set_cell_margins(cell, 100, 140, 100, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "AUTHOR ACTION. ", bold=True, color=RED, size=9)
    add_run(p, text, italic=True, color=GREY, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc: Document, image: Path, caption: str, width_cm: float = 16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    run.add_picture(str(image), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.space_after = Pt(7)
    add_run(cap, caption, size=9)
    return p


def add_placeholder(doc: Document, label: str, instruction: str, caption: str):
    table = doc.add_table(rows=1, cols=1)
    prevent_row_split(table.rows[0])
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(15.8)
    cell = table.cell(0, 0)
    cell.width = Cm(15.8)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(cell, LIGHT_GREY)
    set_cell_margins(cell, 260, 200, 260, 200)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, label, bold=True, color=BLUE, size=11)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p2, instruction, italic=True, color=GREY, size=9)
    cap = doc.add_paragraph()
    cap.paragraph_format.space_after = Pt(7)
    add_run(cap, caption, size=9)


def set_table_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)
            set_cell_margins(row.cells[idx])
            row.cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table(table, header_fill: str = "D9E2F3") -> None:
    set_repeat_table_header(table.rows[0])
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_shading(cell, header_fill if i == 0 else WHITE)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8.5)
                    run.bold = i == 0


def add_source_table(doc: Document) -> None:
    p = doc.add_paragraph()
    add_run(p, "Table 1. Open data sources available to XRD Phase Finder.", bold=True, size=9)
    rows = [
        ("Source", "Primary content", "Role in the program", "Important interpretation"),
        ("COD [2,3]", "Published experimental crystal structures", "Structure retrieval and calculated PXRD patterns", "Open structure collection; calculated patterns are not certified reference measurements"),
        ("Materials Project [14]", "First-principles structures", "Candidate generation and calculated patterns", "Computational structures; agreement must be verified against experiment"),
        ("AFLOW [15]", "First-principles structures", "Candidate generation and calculated patterns", "Computational structures"),
        ("OQMD [16]", "First-principles structures", "Candidate generation and calculated patterns", "Computational structures"),
        ("RRUFF [17]", "Experimental mineral data and structures", "Mineral-oriented reference evidence", "Coverage is domain-specific"),
        ("User CIF", "Any valid local CIF supplied by the user", "Direct comparison and private/reference libraries", "Allows phases absent from integrated sources to be tested"),
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    set_table_widths(table, [2.2, 4.0, 4.2, 5.4])
    style_table(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_validation_table(doc: Document) -> None:
    p = doc.add_paragraph()
    add_run(p, "Table 2. Minimum validation matrix recommended before submission.", bold=True, size=9)
    rows = [
        ("Benchmark block", "Suggested material", "Primary metrics", "Purpose"),
        ("Single phase", "Open experimental RRUFF patterns", "Top-1, top-5, reciprocal rank", "Tests candidate retrieval without mixture ambiguity"),
        ("Synthetic mixtures", "Two- to five-phase profiles generated from held-out structures", "Phase recall, precision, false additions", "Controlled overlap, background, noise and broadening"),
        ("Laboratory patterns", "Expert-labelled minerals, catalysts and solid solutions", "Phase recovery, unexplained-peak fraction", "Tests real backgrounds, preferred orientation and imperfect structures"),
        ("Ablation", "Same patterns under Match-only and Match+Gain ranking", "Change in minor-phase recall and false-positive rate", "Demonstrates the added value of residual scoring"),
        ("Performance", "Cold and warm runs at increasing candidate counts", "Wall time, memory, cache hit rate", "Documents scalability and local-cache benefit"),
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    set_table_widths(table, [2.8, 4.5, 4.0, 4.5])
    style_table(table, header_fill="E2F0D9")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_reference(doc: Document, index: int, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.6)
    p.paragraph_format.space_after = Pt(3)
    add_run(p, f"[{index}] ", size=9)
    add_text_with_links(p, text, size=9)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Heading 1", 13), ("Heading 2", 11.5), ("Heading 3", 10.5)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "XRD Phase Finder — working manuscript", color=GREY, size=8)


def build() -> Path:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(7)
    add_run(
        title,
        "XRD Phase Finder: iterative residual scoring for powder diffraction identification",
        bold=True,
        color=BLUE,
        size=17,
    )
    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(authors, "A. B. Kuznetsov", bold=True, size=11)
    affiliation = doc.add_paragraph()
    affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(
        affiliation,
        "Sobolev Institute of Geology and Mineralogy, Siberian Branch of the Russian Academy of Sciences, "
        "Novosibirsk 630090, Russian Federation",
        italic=True,
        size=9.5,
    )
    email = doc.add_paragraph()
    email.alignment = WD_ALIGN_PARAGRAPH.CENTER
    email.paragraph_format.space_after = Pt(10)
    add_run(email, "Correspondence: ku.artemy@igm.nsc.ru", size=9.5)
    set_paragraph_border(email)

    synopsis = doc.add_paragraph()
    synopsis.paragraph_format.space_before = Pt(8)
    synopsis.paragraph_format.space_after = Pt(5)
    add_run(synopsis, "Synopsis. ", bold=True, color=BLUE)
    add_run(
        synopsis,
        "An open desktop workflow combines multiple open structure sources with global Match ranking and "
        "residual Gain scoring to support inspectable, user-guided identification of crystalline phases.",
        italic=True,
    )

    add_heading(doc, "Abstract", 1)
    add_body(
        doc,
        "Open crystal-structure repositories have made calculated powder diffraction libraries broadly "
        "accessible, but the interpretation of multiphase patterns still requires a transparent distinction "
        "between a candidate that resembles the complete pattern and one that explains intensity left after "
        "previously accepted phases. XRD Phase Finder is an open-source desktop application for preliminary "
        "phase identification from powder X-ray diffraction data. It integrates the Crystallography Open "
        "Database, Materials Project, AFLOW, OQMD and RRUFF resources, while also accepting user-supplied CIF "
        "files. Candidate patterns are calculated locally and inspected together with the experimental profile, "
        "phase markers and unexplained peaks. Initial candidates are ranked by a Match score that combines "
        "strong-peak position coverage, observed-peak coverage and relative-intensity agreement after a bounded "
        "global shift and scale search. Once phases have been selected, a Gain score ranks each remaining "
        "candidate by the weighted residual area that it can explain and penalizes calculated intensity unsupported "
        "by the residual. The graphical workflow keeps database filtering, preprocessing, candidate selection, "
        "profile comparison and project persistence under user control. The software is distributed under the MIT "
        "licence with standalone Windows and macOS packages. Match and Gain are interpretable ranking scores rather "
        "than phase probabilities or quantitative fractions; final structural and quantitative conclusions remain "
        "the domain of profile refinement and expert assessment."
    )
    keywords = doc.add_paragraph()
    keywords.paragraph_format.space_after = Pt(9)
    add_run(keywords, "Keywords: ", bold=True, color=BLUE)
    add_run(
        keywords,
        "powder X-ray diffraction; phase identification; search–match; residual scoring; open crystallographic "
        "databases; graphical user interface; multiphase analysis",
        italic=True,
    )

    add_heading(doc, "1. Introduction", 1)
    add_body(
        doc,
        "Powder X-ray diffraction (PXRD) remains one of the most widely used experimental techniques for phase "
        "identification and crystal-structure characterization because sample preparation is comparatively simple "
        "and the method is applicable to a broad range of crystalline materials [1]. Beyond phase identification, "
        "PXRD is routinely used to determine lattice parameters, crystallite size, microstrain, preferred orientation "
        "and quantitative phase composition. Real laboratory patterns nevertheless remain difficult to interpret: "
        "peaks overlap, relative intensities are affected by preferred orientation and microabsorption, lattice "
        "parameters vary with composition, poorly crystalline components broaden, and instrumental or amorphous "
        "contributions distort the slowly varying signal."
    )
    add_body(
        doc,
        "Software used with powder diffraction data can be divided broadly into search–match systems, profile and "
        "structure-refinement packages, and crystal-structure visualization tools. Search–match applications include "
        "the commercial programs Match! [9], HighScore [10], JADE [11], DIFFRAC.EVA [12] and STOE SEARCH/MATCH 2 "
        "[13]. Refinement-oriented packages include GSAS-II [18], FullProf [19] and TOPAS [20], whereas VESTA [21] "
        "and Mercury [22] are primarily used for structure visualization and crystallographic inspection. These "
        "categories overlap, but their principal tasks remain different: rapid candidate identification, quantitative "
        "profile modelling, or visualization."
    )
    add_body(
        doc,
        "Open crystallographic resources have substantially lowered the barrier to reproducible phase-identification "
        "workflows. The Crystallography Open Database (COD) provides published experimental crystal structures [2,3]. "
        "Materials Project [14], AFLOW [15] and OQMD [16] provide large collections of computational structures, "
        "whereas RRUFF supplies mineral-oriented experimental data and crystal structures [17]. Open libraries "
        "including Gemmi [23], pymatgen [24], spglib [25] and the Atomic Simulation Environment [26] provide standard "
        "operations for reading CIF files, handling structures and applying crystallographic symmetry."
    )
    add_body(
        doc,
        "Several existing programs already connect these resources to phase identification. QUALX2.0 combines "
        "peak-based search–match with the freely available POW_COD database and automated preprocessing [4]. iQual "
        "uses an indexed COD-derived database and a two-stage search strategy [5]. Full-profile search–match tests "
        "candidate structures by iterative Rietveld fitting [6], and FullProfAPP provides a graphical automation "
        "layer for this workflow [7]. XMatcher provides an open framework based on chemical constraints, angular-shift "
        "correction, one-to-one peak matching and peak-level evidence [8]. Commercial systems also support residual "
        "or iterative workflows: notably, DIFFRAC.EVA 8 introduces Profile Fit Residual Search, in which accepted "
        "phases are fitted and unexplained residual regions are searched again [12]. Therefore, neither access to "
        "COD nor residual-guided phase search alone should be presented as unique."
    )
    add_body(
        doc,
        "Despite this progress, an open workflow may still require separate database searches, downloaded structures, "
        "pattern calculations and comparison tools. XRD Phase Finder was developed as a standalone, cross-platform "
        "workspace that combines multiple open structure sources, direct user-supplied CIF input, local indexing, "
        "interactive profile comparison and persistent projects. Its ranking separates two questions. Match estimates "
        "whether a candidate is globally compatible with the experimental peak set. Gain estimates the additional "
        "residual evidence contributed by a candidate after phases already accepted by the user have been fitted. "
        "The distinction is intended to make multiphase interpretation inspectable rather than to replace expert "
        "assessment or full-profile refinement."
    )
    add_body(
        doc,
        "This article describes the software architecture, diffraction calculation, Match and Gain algorithms, and "
        "an evaluation framework for reproducible assessment. XRD Phase Finder is positioned as a "
        "hypothesis-generation and interpretation tool for the stage preceding Rietveld or Le Bail analysis."
    )

    add_heading(doc, "2. Materials and methods", 1)
    add_heading(doc, "2.1. Software architecture and distribution", 2)
    add_body(
        doc,
        "XRD Phase Finder version 1.1.4 is implemented in Python 3.11–3.12. The graphical interface uses PySide6 "
        "[31]; numerical operations use NumPy [27] and SciPy [28]; tabular data handling uses pandas [29]; plotting "
        "uses Matplotlib [30]; and crystallographic information is read with Gemmi [23]. The program "
        "separates data models, file input, crystallographic services, candidate retrieval, scoring and graphical "
        "controllers. Computational modules contain no Qt dependencies. A CalculationContext stores the wavelength, "
        "profile width, zero shift, cell scale and calculation grid as an immutable description of a calculation. "
        "This context is also used to form reproducible cache keys."
    )
    add_body(
        doc,
        "Calculated reflection lists and profiles are cached in memory with bounded least-recently-used behaviour, "
        "and local database metadata and derived peak records are indexed in SQLite. Repeated previews therefore "
        "reuse CIF-to-reflection and reflection-to-profile results when the relevant calculation context has not "
        "changed. Standalone installers are provided separately for Windows and macOS. User settings, project files "
        "and database caches are stored outside the installed application so that program updates do not replace "
        "user data."
    )
    add_figure(
        doc,
        ASSET_DIR / "fig1_workflow.png",
        "Figure 1. Operational workflow of XRD Phase Finder. Experimental data and optional chemical constraints "
        "define a candidate pool assembled from open resources or user CIF files. Match ranks global agreement. "
        "Accepted phases are fitted non-negatively, and Gain re-ranks remaining candidates against the residual. "
        "At every stage the user can inspect and revise the phase model.",
        width_cm=16.0,
    )

    add_heading(doc, "2.2. Input data, preprocessing and analysis state", 2)
    add_body(
        doc,
        "Experimental patterns are imported from common two-column text representations containing diffraction "
        "angle and intensity. The original values are retained. Smoothing and background handling are explicit, "
        "reversible analysis operations rather than hidden destructive transformations. The interface can display "
        "the original profile, a smoothed profile, an estimated physical background, and a combined background plus "
        "broad amorphous contribution. The latter two components can be inspected before subtraction. This distinction "
        "is important because diffuse scattering and broad nanocrystalline peaks are sample information, whereas "
        "instrumental and holder contributions are nuisance signals."
    )
    add_body(
        doc,
        "Automatic settings select parameters for existing smoothing and background models; they do not replace the "
        "numerical algorithms with an opaque classifier. The user may reject the proposed processing and adjust the "
        "parameters manually. For scoring, the program uses the active processed representation and records whether "
        "background removal has been applied. Project files preserve imported patterns, their ordering, selected "
        "phases, processing settings, view settings and relevant phase-assignment state."
    )
    add_body(
        doc,
        "The program provides extended smoothing and background-estimation options. Smoothing includes moving-average "
        "and Gaussian filters and the Savitzky–Golay method [34]. Available background models include asymmetric least "
        "squares (AsLS) [35], asymmetrically reweighted penalized least squares (arPLS) [36], SNIP [37], rolling-ball "
        "processing through pybaselines [38], and robust one-, two- or three-exponential models selected with the "
        "Bayesian information criterion [39]. The background can be estimated alone or together with a broad "
        "amorphous-phase contribution. The physical background and the combined background-plus-amorphous curve are "
        "shown separately; their difference provides an estimate of the broad amorphous contribution. Each component "
        "can be inspected, applied or subtracted under user control."
    )
    add_placeholder(
        doc,
        "INSERT BACKGROUND-DECOMPOSITION EXAMPLE",
        "Use a representative experimental pattern. Show the measured profile in black, the estimated physical "
        "background as a dark thin line, and the combined background-plus-amorphous curve in blue. Shade or label the "
        "difference between the two fitted curves as the broad amorphous contribution. Keep crystalline peak markers "
        "out of this panel so that the decomposition remains immediately readable.",
        "Figure 2. Inspectable separation of slowly varying background and broad amorphous scattering. The physical "
        "background and the combined background-plus-amorphous estimate are displayed independently; the area between "
        "them represents the estimated broad amorphous contribution. No component is subtracted until the user applies "
        "the selected operation.",
    )

    add_heading(doc, "2.3. Open data sources and direct CIF input", 2)
    add_body(
        doc,
        "The candidate layer deliberately distinguishes experimental reference evidence from structures used to "
        "calculate a theoretical pattern. COD contains published experimental crystal structures [2,3]. Materials "
        "Project, AFLOW and OQMD primarily contribute computational structures [14–16]. RRUFF contributes "
        "mineral-oriented experimental data and structures [17]. A valid CIF can also be added directly, which "
        "allows a laboratory reference, a newly refined structure or a phase absent from the integrated services "
        "to be evaluated without modifying a global database."
    )
    add_source_table(doc)
    add_body(
        doc,
        "Chemical-element filters restrict the initial candidate pool but are not treated as proof of phase presence. "
        "Source and entry identifiers are retained together to prevent collisions between databases. Each displayed "
        "candidate links its ranking evidence to a traceable source record or local CIF."
    )

    add_heading(doc, "2.4. Calculation of candidate patterns", 2)
    add_body(
        doc,
        "For a structure-based candidate, symmetry-expanded atomic positions, unit-cell parameters and the selected "
        "wavelength are used to calculate reflection positions and relative intensities. The calculation follows "
        "standard powder-diffraction relations: Bragg positions, structure-factor summation with Cromer–Mann atomic "
        "scattering factors [40], multiplicity and displacement terms, and a Lorentz–polarization correction consistent "
        "with standard tabulations [41]. Reflection sticks can be broadened into a pseudo-Voigt profile [42]. Profile "
        "contributions are evaluated only in a finite interval around each reflection, avoiding full-grid evaluation "
        "for points where the contribution is negligible. Kα2 can be included when appropriate. Candidate sticks, "
        "broadened profiles and the experimental signal are displayed on a shared angular scale; optional global zero "
        "and small scale corrections are used for ranking but are reported as alignment aids, not as refined "
        "crystallographic parameters."
    )

    add_heading(doc, "2.5. Observed peak detection", 2)
    add_body(
        doc,
        "Observed peaks are detected with a prominence threshold that combines a robust noise estimate and a fraction "
        "of the 95th percentile of the active intensity signal. Peak prominence and width are evaluated with the SciPy "
        "signal-processing implementation [28], while noise is estimated from the median absolute deviation of first "
        "differences, following the robust-statistics principle described by Hampel [43]. Candidate ranking uses up to "
        "80 observed records, while the strongest 24 calculated reflections form the main candidate-side evidence. "
        "The minimum peak separation corresponds to approximately 0.11° 2θ and the accepted detected width is bounded "
        "to suppress isolated noise and excessively broad baseline features. These defaults are pragmatic search "
        "settings and should be reported with benchmark results."
    )

    add_heading(doc, "2.6. Match: global candidate compatibility", 2)
    add_body(
        doc,
        "For each candidate, the program searches a bounded alignment grid consisting of additive zero shifts and a "
        "small angular scale deformation around a 45° pivot. The current implementation evaluates zero shifts between "
        "−0.35 and +0.35° 2θ, supplemented by a robust shift estimate, and scale terms between −0.3 and +0.3%. "
        "Calculated and observed peaks are paired within a tolerance of 0.34° 2θ. The strongest calculated reflections "
        "receive weights proportional to relative intensity raised to 0.55; observed peaks receive a weaker intensity "
        "weight so that unrelated reflections from other phases do not dominate a candidate's score."
    )
    add_equation(
        doc,
        "M₀ = 100 [0.78 Ccalc + 0.16 Cobs + 0.06 min(Ntop / 6, 1)]",
    )
    add_body(
        doc,
        "Here Ccalc is weighted coverage of the candidate reflections, Cobs is weighted coverage of the strongest "
        "observed peaks, and Ntop is the number of matched reflections among the eight strongest calculated lines. "
        "The preliminary score M0 is multiplied by a bounded deformation penalty and by 0.72 + 0.28 QI, where QI "
        "measures relative-intensity agreement after a least-squares scale. Additional caps prevent a high score when "
        "the candidate's strongest line or too few of its leading reflections are present. The displayed Match value "
        "is therefore an interpretable heuristic rank in the interval 0–100%, not a posterior probability."
    )

    add_heading(doc, "2.7. Gain: conditional residual contribution", 2)
    add_body(
        doc,
        "Gain is calculated only after at least one phase has been selected. Profiles of all selected phases are fitted "
        "simultaneously to the active non-negative target by weighted non-negative least squares using the "
        "Lawson–Hanson formulation [44]. Their summed profile is subtracted, and the positive residual is lightly "
        "smoothed at a scale related to the estimated peak width. For each remaining candidate, pseudo-Voigt profiles "
        "are tested over three profile-width multipliers and several mixing parameters. A non-negative scale factor is "
        "fitted to the residual."
    )
    add_body(
        doc,
        "Let R be the total weighted residual area, C the weighted area covered by both the residual and the scaled "
        "candidate profile, and E the weighted excess area predicted by the candidate where the residual provides no "
        "support. The implemented score is"
    )
    add_equation(doc, "G = 100 (C / R) [C / (C + 3E)].")
    add_body(
        doc,
        "The first factor rewards the fraction of residual evidence explained; the second penalizes unsupported "
        "calculated intensity three times more strongly than covered area. A candidate duplicating an already selected "
        "phase therefore tends to receive low Gain even if its global Match remains high. Conversely, a weaker phase "
        "can move upward if its reflections coincide with previously unexplained residual features. Gain is a "
        "state-dependent ranking score, not a concentration estimate."
    )
    add_figure(
        doc,
        ASSET_DIR / "fig4_match_gain.png",
        "Figure 3. Conceptual distinction between Match and Gain. (a) Match compares a candidate with the complete "
        "observed peak set. (b) Selected phases define explained intensity and a positive residual. (c) Gain is low for "
        "a duplicate candidate and higher for a candidate that covers residual evidence without predicting substantial "
        "unsupported intensity. The equation is the expression implemented in version 1.1.4.",
        width_cm=16.0,
    )

    add_heading(doc, "2.8. Reproducibility and software availability", 2)
    add_body(
        doc,
        "Source code and tagged releases are available at https://github.com/ABKuznetsov/XRD_Analysis_Toolkit under "
        "the MIT licence. The repository contains the application code, release-building scripts and platform-specific "
        "packaging. A release used in a publication should be archived with a persistent DOI, and benchmark input "
        "patterns, expected phase labels and analysis settings should be deposited separately under an open licence. "
        "Database records remain subject to the terms and citations of their original providers."
    )

    add_heading(doc, "3. Evaluation design", 1)
    add_heading(doc, "3.1. Position relative to existing software", 2)
    add_body(
        doc,
        "The relevant comparison is not whether a program can display COD structures. QUALX2.0 [4] and iQual [5] "
        "already provide efficient peak-based search against COD-derived collections. Full-profile search–match [6] "
        "and FullProfAPP [7] provide more rigorous profile-based testing and automated refinement, while XMatcher [8] "
        "is a close open and evidence-oriented comparison. Among commercial systems, Match! supports both peak-based "
        "and profile-fitting search–match and distributes a COD-derived reference database [9]; HighScore supports "
        "multiple reference databases and integrates phase identification with profile and Rietveld analysis [10]; "
        "JADE combines phase identification and whole-pattern fitting [11]; DIFFRAC.EVA 8 performs iterative Profile "
        "Fit Residual Search [12]; and STOE SEARCH/MATCH 2 supports COD together with background correction, peak "
        "processing and profile fitting [13]."
    )
    add_body(
        doc,
        "XRD Phase Finder should therefore be evaluated on the combination that defines its intended niche rather "
        "than on any one feature in isolation: open-source distribution, candidate generation from several open "
        "experimental and computational repositories [2,3,14–17], direct CIF inclusion, transparent peak-level and "
        "residual evidence, persistent interactive projects, and the ability of Gain to prioritize an additional "
        "candidate after accepted phases have been fitted. This claim must be tested directly against Match-only "
        "ranking and, where licences permit, against representative open and commercial search–match workflows "
        "[4–13]."
    )
    add_body(
        doc,
        "A fair benchmark should use the same chemical constraints and candidate availability for every method. "
        "Database coverage failures must be separated from ranking failures. Calculated structures should also be "
        "reported separately from experimental reference patterns because their relative intensities and relaxed "
        "lattice parameters have different evidential status."
    )

    add_heading(doc, "3.2. Validation datasets and metrics", 2)
    add_validation_table(doc)
    add_author_action(
        doc,
        "Populate Table 2 with actual sample counts and measured results. For a Journal of Applied Crystallography "
        "submission, include at least two independent users, public benchmark patterns, top-k accuracy, multiphase "
        "precision/recall, runtime and a Match-only versus Match+Gain ablation. Do not replace this evidence with "
        "screenshots or selected success cases."
    )

    add_heading(doc, "4. Results and discussion", 1)
    add_heading(doc, "4.1. Interactive evidence workspace", 2)
    add_body(
        doc,
        "The principal application window combines an experimental pattern, candidate table, element filters, selected "
        "phase list, structure card and database controls. The plot can show observed and processed signals, calculated "
        "profiles, reflection markers, fitted total intensity, background components and unexplained peaks. Candidate "
        "selection updates the preview without requiring a new database search. Resizable panels and persistent view "
        "settings allow the same workspace to be used on compact and large displays."
    )
    add_placeholder(
        doc,
        "INSERT APPLICATION SCREENSHOT",
        "Use one clean 16:9 screenshot with an experimental multiphase pattern, two or three selected phases, a visible "
        "candidate table and unexplained-peak markers. Crop only operating-system chrome, not scientific content.",
        "Figure 4. Main XRD Phase Finder workspace. The central plot combines experimental and calculated evidence; "
        "candidate and selected-phase tables remain visible so that ranking decisions can be inspected in context.",
    )

    add_heading(doc, "4.2. Dominant-phase ranking", 2)
    add_body(
        doc,
        "In the first iteration no phase model exists, so candidates are ordered primarily by Match. This stage is "
        "designed to recover a dominant phase even when additional observed peaks belong to impurities. Candidate-side "
        "coverage is therefore weighted more strongly than observed-side coverage. Visual acceptance should require "
        "agreement of several leading reflections and inspection of strong predicted lines that are absent from the "
        "experiment; a single coincident maximum is insufficient."
    )
    add_placeholder(
        doc,
        "INSERT DOMINANT-PHASE EXAMPLE",
        "Recommended example: natural gehlenite. Show the raw or conservatively processed profile, the selected "
        "gehlenite candidate, calculated markers and the remaining unexplained peaks. Include COD/source identifiers.",
        "Figure 5. Dominant-phase identification example. Match ranks the phase against the complete observed peak set, "
        "while absent calculated peaks and remaining observed peaks remain visible for expert assessment.",
    )

    add_heading(doc, "4.3. Residual ranking in multiphase patterns", 2)
    add_body(
        doc,
        "After the dominant phase is accepted, candidate ordering changes from global similarity to incremental "
        "explanatory value. This distinction is most useful when several database entries represent the same structure "
        "type or when a minor phase shares its strongest reflection with the dominant phase. A duplicate can retain a "
        "high Match because it resembles the complete pattern, yet receive a low Gain because the selected model has "
        "already covered the corresponding intensity. A candidate with a lower Match can receive higher Gain when it "
        "accounts for residual peaks without introducing intense unsupported reflections."
    )
    add_placeholder(
        doc,
        "INSERT MATCH–GAIN ABLATION EXAMPLE",
        "Use the same pattern twice: left, candidates ordered by Match after the first phase; right, ordered by Gain. "
        "Mark the first correctly added minor phase and at least one high-Match duplicate demoted by Gain.",
        "Figure 6. Effect of conditional residual ranking. Comparison of Match-only and Match+Gain ordering after "
        "selection of the dominant phase demonstrates whether Gain promotes an additional supported phase rather than "
        "a duplicate explanation.",
    )

    add_heading(doc, "4.4. Background, amorphous scattering and broad peaks", 2)
    add_body(
        doc,
        "Background treatment is a dominant source of uncertainty for catalyst and mineral patterns. A flexible "
        "baseline can incorrectly follow the lower shoulders of crystalline peaks, whereas an overly rigid exponential "
        "background can leave holder scattering or diffuse intensity in the residual. XRD Phase Finder therefore "
        "separates an estimated physical background from an optional broad amorphous contribution and allows both to "
        "be visualized before subtraction. This makes the analysis decision inspectable and preserves the possibility "
        "of showing an amorphous halo on the original experimental profile."
    )
    add_body(
        doc,
        "The present automatic mode selects parameters among explicit smoothing and baseline models. It should not be "
        "interpreted as a universal physical decomposition of instrumental background, amorphous material and "
        "nanocrystalline broadening. Validation must include patterns with low-angle holder contributions, diffuse "
        "halos, sharp and broad crystalline peaks, negative detector values and overlapping phases. Performance should "
        "be assessed by peak-area preservation and by stability of phase ranking, not by visual flatness alone."
    )

    add_heading(doc, "4.5. Limitations", 2)
    add_body(
        doc,
        "Phase identification cannot succeed when the relevant structure or reference pattern is absent from the "
        "candidate pool. Closely related polymorphs, solid solutions and structures with similar unit cells may remain "
        "ambiguous. Calculated intensities can disagree with experiment because of preferred orientation, absorption, "
        "texture, site occupancy, radiation choice and incomplete structural models. Bounded zero and scale corrections "
        "improve robust ranking but do not constitute lattice refinement. Match and Gain should therefore guide "
        "inspection; they must not be reported as statistical confidence or mass percentage."
    )
    add_body(
        doc,
        "Quantitative phase analysis requires appropriate profile and instrument models, scale factors, absorption "
        "treatment and, where relevant, an internal standard or reference-intensity-ratio strategy. XRD Phase Finder "
        "can supply a candidate phase set and an auditable interpretation state for such subsequent analysis, but the "
        "current version does not replace a validated Rietveld or Le Bail workflow."
    )

    add_heading(doc, "5. Conclusions", 1)
    add_body(
        doc,
        "XRD Phase Finder provides a cross-platform, open-source environment for preliminary phase identification from "
        "powder diffraction data. Its contribution is not merely access to COD, which is already available through "
        "several mature programs. The program instead combines multiple open structure sources and direct CIF input "
        "with an inspectable iterative workflow. Match ranks initial global compatibility; Gain ranks the additional "
        "residual evidence supplied by a candidate after accepted phases have been fitted. The interface preserves "
        "supporting and conflicting peak evidence, processing choices and project state."
    )
    add_body(
        doc,
        "The strongest route to a high-impact software publication is a public benchmark demonstrating when residual "
        "Gain improves minor-phase recovery without increasing false assignments, together with runtime measurements "
        "and independent-user evaluation. Such validation would establish XRD Phase Finder as a reproducible bridge "
        "between rapid open-database screening and full-profile refinement."
    )

    add_heading(doc, "Acknowledgements", 1)
    add_body(
        doc,
        "This work was supported by the state assignment of the Sobolev Institute of Geology and Mineralogy, Siberian "
        "Branch of the Russian Academy of Sciences (No. FWZN-2026-0014)."
    )
    add_heading(doc, "Author contributions", 1)
    add_body(
        doc,
        "A. B. Kuznetsov: conceptualization, methodology, software, validation, investigation, visualization, writing "
        "and project administration."
    )
    add_heading(doc, "Conflict of interest", 1)
    add_body(doc, "The author declares no conflict of interest.")
    add_heading(doc, "Data and code availability", 1)
    add_body(
        doc,
        "The source code and release packages are available at "
        "https://github.com/ABKuznetsov/XRD_Analysis_Toolkit. The exact release used for the reported evaluation, "
        "benchmark patterns, ground-truth labels and analysis settings will be archived in a persistent repository "
        "before submission."
    )
    add_author_action(
        doc,
        "Create a Zenodo release DOI and replace the future tense above. Deposit only data that can legally be "
        "redistributed. Add an explicit disclosure of any generative-AI assistance in accordance with the target "
        "journal's current policy."
    )

    add_heading(doc, "References", 1)
    references = [
        "Dinnebier, R. E. & Billinge, S. J. L. (eds.) (2008). Powder Diffraction: Theory and Practice. Cambridge: Royal Society of Chemistry.",
        "Gražulis, S. et al. (2009). Crystallography Open Database — an open-access collection of crystal structures. J. Appl. Cryst. 42, 726–729. https://doi.org/10.1107/S0021889809016690.",
        "Gražulis, S. et al. (2012). Crystallography Open Database (COD): an open-access collection of crystal structures and platform for world-wide collaboration. Nucleic Acids Res. 40, D420–D427. https://doi.org/10.1093/nar/gkr900. Resource: https://www.crystallography.net/cod/.",
        "Altomare, A., Corriero, N., Cuocci, C., Falcicchio, A., Moliterni, A. & Rizzi, R. (2015). QUALX2.0: a qualitative phase analysis software using the freely available database POW_COD. J. Appl. Cryst. 48, 598–603. https://doi.org/10.1107/S1600576715002319. Software: https://www.ba.ic.cnr.it/softwareic/qualxweb/.",
        "Lv, B., Feng, Z., Sun, X., Cao, J., Gao, Y., Chang, S., Dong, C. & Zhang, J. (2024). iQual: a new computer program for qualitative phase analysis with efficient search–match capability. J. Appl. Cryst. 57, 572–579. https://doi.org/10.1107/S1600576724002267.",
        "Lutterotti, L., Pillière, H., Fontugne, C., Boullay, P. & Chateigner, D. (2019). Full-profile search–match by the Rietveld method. J. Appl. Cryst. 52, 587–598. https://doi.org/10.1107/S160057671900342X.",
        "Arcelus, O. et al. (2024). FullProfAPP: a graphical user interface for the streamlined automation of powder diffraction data analysis. J. Appl. Cryst. 57, 1676–1690. https://doi.org/10.1107/S1600576724006885.",
        "Cao, B. (2026). XMatcher: an open-source framework for X-ray diffraction phase identification. arXiv:2607.17162. https://arxiv.org/abs/2607.17162.",
        "Putz, H. & Brandenburg, K. (2024). Match! — Phase Analysis using Powder Diffraction. Crystal Impact, Bonn, Germany. https://www.crystalimpact.com/match/.",
        "Degen, T., Sadki, M., Bron, E., König, U. & Nénert, G. (2014). The HighScore suite. Powder Diffr. 29(S2), S13–S18. https://doi.org/10.1017/S0885715614000840.",
        "International Centre for Diffraction Data (2026). JADE XRD software: phase identification and whole-pattern fitting. https://www.icdd.com/mdi-jade/ (accessed 25 July 2026).",
        "Bruker (2026). DIFFRAC.EVA 8: XRD evaluation software and Profile Fit Residual Search. https://www.bruker.com/en/products-and-solutions/diffractometers-and-x-ray-microscopes/x-ray-diffractometers/diffrac-suite-software/diffrac-eva.html (accessed 25 July 2026).",
        "STOE & Cie GmbH (2026). The new STOE SEARCH/MATCH 2. https://www.stoe.com/articles/new-search-match-2-appnote/ (accessed 25 July 2026).",
        "Jain, A. et al. (2013). Commentary: The Materials Project: a materials genome approach to accelerating materials innovation. APL Mater. 1, 011002. https://doi.org/10.1063/1.4812323. Resource: https://materialsproject.org/.",
        "Curtarolo, S. et al. (2012). AFLOW: an automatic framework for high-throughput materials discovery. Comput. Mater. Sci. 58, 218–226. https://doi.org/10.1016/j.commatsci.2012.02.005. Resource: https://aflow.org/.",
        "Saal, J. E., Kirklin, S., Aykol, M., Meredig, B. & Wolverton, C. (2013). Materials design and discovery with high-throughput density functional theory: the Open Quantum Materials Database (OQMD). JOM 65, 1501–1509. https://doi.org/10.1007/s11837-013-0755-4. Resource: https://oqmd.org/.",
        "Lafuente, B., Downs, R. T., Yang, H. & Stone, N. (2015). The power of databases: the RRUFF project. In Highlights in Mineralogical Crystallography, edited by T. Armbruster & R. M. Danisi, pp. 1–30. Berlin: De Gruyter. https://doi.org/10.1515/9783110417104-003. Resource: https://rruff.info/.",
        "Toby, B. H. & Von Dreele, R. B. (2013). GSAS-II: the genesis of a modern open-source all-purpose crystallography software package. J. Appl. Cryst. 46, 544–549. https://doi.org/10.1107/S0021889813003531.",
        "Rodríguez-Carvajal, J. (1993). Recent advances in magnetic structure determination by neutron powder diffraction. Physica B 192, 55–69. https://doi.org/10.1016/0921-4526(93)90108-I.",
        "Coelho, A. A. (2018). TOPAS and TOPAS-Academic: an optimization program integrating computer algebra and crystallographic objects written in C++. J. Appl. Cryst. 51, 210–218. https://doi.org/10.1107/S1600576718000183.",
        "Momma, K. & Izumi, F. (2011). VESTA 3 for three-dimensional visualization of crystal, volumetric and morphology data. J. Appl. Cryst. 44, 1272–1276. https://doi.org/10.1107/S0021889811038970.",
        "Macrae, C. F. et al. (2008). Mercury CSD 2.0 — new features for the visualization and investigation of crystal structures. J. Appl. Cryst. 41, 466–470. https://doi.org/10.1107/S0021889807067908.",
        "Wojdyr, M. (2022). GEMMI: a library for structural biology. J. Open Source Softw. 7, 4200. https://doi.org/10.21105/joss.04200.",
        "Ong, S. P. et al. (2013). Python Materials Genomics (pymatgen): a robust, open-source Python library for materials analysis. Comput. Mater. Sci. 68, 314–319. https://doi.org/10.1016/j.commatsci.2012.10.028.",
        "Togo, A. & Tanaka, I. (2018). Spglib: a software library for crystal symmetry search. arXiv:1808.01590. https://arxiv.org/abs/1808.01590.",
        "Larsen, A. H. et al. (2017). The Atomic Simulation Environment — a Python library for working with atoms. J. Phys. Condens. Matter 29, 273002. https://doi.org/10.1088/1361-648X/aa680e.",
        "Harris, C. R. et al. (2020). Array programming with NumPy. Nature 585, 357–362. https://doi.org/10.1038/s41586-020-2649-2.",
        "Virtanen, P. et al. (2020). SciPy 1.0: fundamental algorithms for scientific computing in Python. Nat. Methods 17, 261–272. https://doi.org/10.1038/s41592-019-0686-2.",
        "McKinney, W. (2010). Data structures for statistical computing in Python. Proc. 9th Python in Science Conference, 56–61. https://doi.org/10.25080/Majora-92bf1922-00a.",
        "Hunter, J. D. (2007). Matplotlib: a 2D graphics environment. Comput. Sci. Eng. 9, 90–95. https://doi.org/10.1109/MCSE.2007.55.",
        "The Qt Company (2026). Qt for Python (PySide6) documentation. https://doc.qt.io/qtforpython-6/ (accessed 25 July 2026).",
        "Doebelin, N. & Kleeberg, R. (2015). Profex: a graphical user interface for the Rietveld refinement program BGMN. J. Appl. Cryst. 48, 1573–1580. https://doi.org/10.1107/S1600576715014685.",
        "Rietveld, H. M. (1969). A profile refinement method for nuclear and magnetic structures. J. Appl. Cryst. 2, 65–71. https://doi.org/10.1107/S0021889869006558.",
        "Savitzky, A. & Golay, M. J. E. (1964). Smoothing and differentiation of data by simplified least squares procedures. Anal. Chem. 36, 1627–1639. https://doi.org/10.1021/ac60214a047.",
        "Eilers, P. H. C. & Boelens, H. F. M. (2005). Baseline correction with asymmetric least squares smoothing. Leiden University Medical Centre Report.",
        "Baek, S.-J., Park, A., Ahn, Y.-J. & Choo, J. (2015). Baseline correction using asymmetrically reweighted penalized least squares smoothing. Analyst 140, 250–257. https://doi.org/10.1039/C4AN01061B.",
        "Morháč, M., Kliman, J., Matoušek, V., Veselský, M. & Turzo, I. (1997). Background elimination methods for multidimensional coincidence gamma-ray spectra. Nucl. Instrum. Methods Phys. Res. A 401, 113–132. https://doi.org/10.1016/S0168-9002(97)01023-1.",
        "Erb, D. (2022). pybaselines: a Python library of algorithms for the baseline correction of experimental data. https://doi.org/10.5281/zenodo.5608581.",
        "Schwarz, G. (1978). Estimating the dimension of a model. Ann. Stat. 6, 461–464. https://doi.org/10.1214/aos/1176344136.",
        "Cromer, D. T. & Mann, J. B. (1968). X-ray scattering factors computed from numerical Hartree–Fock wave functions. Acta Cryst. A24, 321–324. https://doi.org/10.1107/S0567739468000550.",
        "Gilmore, C. J., Kaduk, J. A. & Schenk, H. (eds.) (2019). International Tables for Crystallography, Vol. H: Powder Diffraction. Chester: International Union of Crystallography. https://doi.org/10.1107/97809553602060000115.",
        "Thompson, P., Cox, D. E. & Hastings, J. B. (1987). Rietveld refinement of Debye–Scherrer synchrotron X-ray data from Al2O3. J. Appl. Cryst. 20, 79–83. https://doi.org/10.1107/S0021889887087090.",
        "Hampel, F. R. (1974). The influence curve and its role in robust estimation. J. Am. Stat. Assoc. 69, 383–393. https://doi.org/10.1080/01621459.1974.10482962.",
        "Lawson, C. L. & Hanson, R. J. (1995). Solving Least Squares Problems. Philadelphia: Society for Industrial and Applied Mathematics. https://doi.org/10.1137/1.9781611971217.",
    ]
    for idx, reference in enumerate(references, 1):
        add_reference(doc, idx, reference)

    doc.core_properties.title = "XRD Phase Finder: iterative residual scoring for powder diffraction identification"
    doc.core_properties.subject = "Working manuscript for Journal of Applied Crystallography"
    doc.core_properties.author = "A. B. Kuznetsov"
    doc.core_properties.keywords = "PXRD, phase identification, search-match, residual scoring, open databases"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
