from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
MATCH_GAIN_FIGURE = ROOT / "manuscript_assets" / "phase_finder" / "fig4_match_gain.png"


REPLACEMENTS = {
    (
        "Open desktop software supports automatic and semi-automatic interpretation of crystalline phases "
        "in powder mixtures using several structure sources, user CIF files, Match ranking and residual Gain scoring."
    ): (
        "Open desktop software supports automatic and semi-automatic interpretation of crystalline phases "
        "in powder mixtures using several structure sources, user CIF files, global Match ranking and staged Gain scoring."
    ),
    (
        "XRD Phase Finder is an open-source desktop application for automatic and semi-automatic interpretation "
        "of the crystalline phases present in powder mixtures. Automatic search generates and ranks candidates, "
        "whereas the semi-automatic workflow allows the user to set elemental constraints, inspect the diffraction "
        "evidence and accept phases between residual-search iterations. The program searches user-accessible records "
        "from the Crystallography Open Database, Materials Project, AFLOW, OQMD and RRUFF, and it can evaluate an "
        "individual CIF file without adding it to a database. Candidate patterns are calculated locally and displayed "
        "with the experimental profile, reflection markers, background components and unexplained lines. Match ranks "
        "initial candidates against the complete set of experimental line records. After one or more phases have been "
        "accepted, Gain re-ranks the remaining candidates using the residual line set and demotes entries that duplicate "
        "already explained reflections. Smoothing, physical-background estimation and an optional broad amorphous "
        "contribution remain inspectable and reversible. For accepted phases, normalized non-negative fitted profile "
        "contributions provide a semi-quantitative estimate. Projects retain imported patterns, processing settings, "
        "accepted phases and view state. Standalone Windows and macOS packages are distributed under the MIT licence. "
        "Match and Gain are ranking scores rather than phase probabilities, and the semi-quantitative values are not "
        "validated mass fractions."
    ): (
        "XRD Phase Finder is an open-source desktop application for automatic and semi-automatic interpretation "
        "of the crystalline phases present in powder mixtures. Automatic search generates and ranks candidates, "
        "whereas the semi-automatic workflow allows the user to set elemental constraints, inspect the diffraction "
        "evidence and accept phases between search iterations. The program searches user-accessible records from the "
        "Crystallography Open Database, Materials Project, AFLOW, OQMD and RRUFF, and it can evaluate an individual "
        "CIF file without adding it to a database. Candidate patterns are calculated locally and displayed with the "
        "experimental profile, reflection markers, background components and unexplained lines. Match ranks initial "
        "candidates by two-way agreement between their strongest reflections and the strongest experimental line "
        "records. After one or more phases have been accepted and jointly fitted, Gain searches successively for "
        "uncovered lines, under-fitted overlapping lines and weaker hidden-phase evidence. It rewards additional "
        "supported diffraction signal while penalizing duplicate or unsupported strong reflections. Smoothing, "
        "physical-background estimation and an optional broad amorphous contribution remain inspectable and reversible. "
        "For accepted phases, normalized non-negative fitted profile contributions provide a semi-quantitative estimate. "
        "Projects retain imported patterns, processing settings, accepted phases and view state. Standalone Windows "
        "and macOS packages are distributed under the MIT licence. Match and Gain are ranking scores rather than phase "
        "probabilities, and the semi-quantitative values are not validated mass fractions."
    ),
    (
        "In practice, an open workflow still often involves separate database searches, downloaded structures, pattern "
        "calculations and comparison tools. XRD Phase Finder brings these steps into one cross-platform workspace. It "
        "searches several open structure sources, accepts user CIF files, builds local indexes, compares profiles "
        "interactively and saves the analysis as a project. The ranking answers two different questions. Match asks "
        "whether a candidate agrees with the experimental line set as a whole. Gain asks what additional residual "
        "evidence remains for that candidate after the user has accepted other phases. This separation keeps the "
        "reasoning in a multiphase search visible to the user."
    ): (
        "In practice, an open workflow still often involves separate database searches, downloaded structures, pattern "
        "calculations and comparison tools. XRD Phase Finder brings these steps into one cross-platform workspace. It "
        "searches several open structure sources, accepts user CIF files, builds local indexes, compares profiles "
        "interactively and saves the analysis as a project. The ranking answers two different questions. Match asks "
        "whether a candidate agrees with the experimental line set as a whole. Gain asks what that candidate adds to "
        "the phase model already accepted by the user, including signal hidden beneath an incompletely fitted shared "
        "reflection. This separation keeps the reasoning in a multiphase search visible to the user."
    ),
    (
        "XRD Phase Finder version 1.1.4 is implemented in Python 3.11–3.12."
    ): (
        "XRD Phase Finder version 1.2.0 is implemented in Python 3.11–3.12."
    ),
    (
        "Operational workflow of XRD Phase Finder. Experimental data and optional chemical constraints define a "
        "candidate pool assembled from open resources or user CIF files. Match ranks global agreement. Accepted phases "
        "are fitted non-negatively, and Gain re-ranks remaining candidates against the residual. At every stage the "
        "user can inspect and revise the phase model."
    ): (
        "Operational workflow of XRD Phase Finder. Experimental data and optional chemical constraints define a "
        "candidate pool assembled from open resources or user CIF files. Match ranks global agreement. Accepted phase "
        "profiles are fitted jointly with non-negative coefficients, and staged Gain re-ranks remaining candidates "
        "using uncovered, overlapping and hidden-phase evidence. At every stage the user can inspect and revise the "
        "phase model."
    ),
    (
        "Match is a global stick-line compatibility score. For each candidate, the calculated reflection list is "
        "treated as a set of stick lines with relative intensities; the candidate is not primarily judged by a broadened "
        "synthetic profile. The experimental pattern is represented by the line records described above. A calculated "
        "line is considered supported when it falls inside the FWHM-controlled window of an experimental line; a central "
        "plateau is used so that small zero-shift or lattice-parameter errors do not reduce the score. The score combines "
        "direct candidate-line coverage with a weaker line-pair fingerprint term based on relative spacings in sin(theta). "
        "The pair term rewards a consistent internal spacing pattern, while the direct line term remains dominant so that "
        "partly overlapped or missing weak lines do not over-penalize a correct phase."
    ): (
        "Match is a global stick-line compatibility score. For each candidate, the strongest calculated reflections "
        "are compared with the strongest experimental line records rather than with every point of a broadened synthetic "
        "profile. A bounded scale in sin(theta) is estimated from clusters of direct line coincidences, allowing small "
        "systematic angular differences without treating them as a refined unit cell. The aligned comparison is "
        "bidirectional: observed-anchor coverage asks whether the candidate explains the strongest experimental lines, "
        "whereas candidate-line coverage asks whether the candidate's own strong reflections are supported by the "
        "experiment. Position agreement is evaluated with a narrow central plateau and a linearly decreasing tolerance, "
        "and relative-intensity agreement moderates each line match."
    ),
    (
        "M = 100 [0.70 Cline + 0.18 Cpair + 0.12 min(Nline / 6, 1)]."
    ): (
        "M = 100 [0.62 Cobs + 0.25 Cref + 0.08 Ccount + 0.05 Cseed]."
    ),
    (
        "Here Cline is the weighted coverage of the candidate's strongest calculated stick lines by experimental line "
        "windows, Cpair is the normalized support of internally consistent line-pair spacings, and Nline is the number "
        "of supported leading lines. Observed-line areas are used as local support weights but are not normalized into "
        "a complete single-phase experimental pattern, because strong peaks from other phases may dominate a multiphase "
        "sample. The displayed Match value is therefore an interpretable heuristic rank in the interval 0-100%, not a "
        "posterior probability or concentration estimate."
    ): (
        "Here Cobs is the intensity-weighted coverage of the leading observed anchors, Cref is the weighted support of "
        "the candidate reflections, Ccount increases as both directions accumulate independent matches, and Cseed "
        "describes the strength of the alignment cluster. Scores are capped when too few observed anchors or candidate "
        "lines are supported, preventing a small number of accidental coincidences from producing a high rank. The "
        "displayed Match value is therefore an interpretable heuristic score in the interval 0–100%, not a posterior "
        "probability or concentration estimate."
    ),
    (
        "Gain is a conditional residual-line score. After one or more phases have been selected, their calculated stick "
        "positions and fitted profile contributions are used to mark experimental lines as explained. A line is removed "
        "from the residual candidate evidence if it lies within the FWHM-controlled window of a selected phase line, or "
        "if the selected total profile covers the integrated line region. The remaining residual line set is then scored "
        "with the same stick-line compatibility function used for Match. When no phase has been selected, Gain is "
        "identical to Match; after selection, a duplicate or near-duplicate phase can retain a high Match but should "
        "receive low Gain because its strongest lines have already been explained."
    ): (
        "Gain is a conditional score for the additional diffraction evidence supplied by a candidate. The profiles of "
        "all accepted phases are first refitted jointly to the active background-corrected signal by weighted "
        "non-negative least squares. This joint step allows an earlier phase coefficient to decrease when a later phase "
        "shares a strong reflection, instead of assigning all shared intensity permanently to the first phase. The "
        "positive, noise-controlled difference between experiment and the selected-phase total defines the current "
        "residual evidence. Gain is displayed only after at least one phase has been accepted and is not calculated by "
        "reusing the Match formula."
    ),
    (
        "The next candidate is not chosen by fitting another broadened calculated profile to the residual. Instead, the "
        "ranking asks whether its stick lines are supported by the remaining experimental line integrals. This keeps the "
        "residual search tied to discrete, visible line evidence."
    ): (
        "The search proceeds through three evidence stages. The direct stage uses experimental lines not explained by "
        "selected sticks or profiles. If this evidence is sparse, the overlap stage examines observed lines close to "
        "selected reflections but still under-fitted in height by at least a local deficit. The hidden stage is used only "
        "when neither of the first two stages supplies enough independent lines; it tests whether a candidate has "
        "coherent support across the complete observed line set despite extensive overlap."
    ),
    (
        "G = F(candidate sticks, residual experimental lines)."
    ): (
        "Candidate sticks provide the primary Gain evidence. Their score increases with coverage of major residual or "
        "under-fitted lines, support among the candidate's leading reflections and consistency of relative line "
        "intensities. It decreases when strong calculated reflections are absent from the experiment or merely repeat "
        "already fitted lines. A broadened candidate profile then moderates, but cannot entirely erase, valid stick-line "
        "support; calculated intensity above the residual is penalized more strongly than remaining under-fit intensity."
    ),
    (
        "The practical effect is a state-dependent ordering: after a dominant phase is accepted, candidates that explain "
        "only already-covered reflections are demoted, whereas a weaker phase can move upward if its stick lines coincide "
        "with residual line integrals. Gain is not a concentration estimate and must be interpreted together with the "
        "plot, source identifiers, I/Ic values and candidate card."
    ): (
        "The resulting Gain is bounded by the fraction of the weighted profile fit still unexplained by the accepted "
        "phase model. Ranking stops when the selected-phase fit or residual-area criteria indicate that another phase "
        "would probably overfit noise. Thus, a duplicate phase can retain a high global Match but receives little or no "
        "Gain, whereas a weaker phase can move upward through uncovered or demonstrably under-fitted reflections. Gain "
        "is not a concentration estimate and must be interpreted together with the difference curve, source identifiers, "
        "I/Ic values and candidate card."
    ),
    (
        "Conceptual distinction between Match and Gain. (a) The observed profile is converted into experimental line "
        "records with position, integrated area and FWHM. (b) Match compares a candidate's calculated stick lines with "
        "the complete observed line set. (c) Selected phases mark covered experimental lines. (d) Gain compares a "
        "remaining candidate only with the residual line set, reducing the score of duplicate explanations while "
        "promoting phases that cover previously unexplained line integrals."
    ): (
        "Current Match and Gain logic. (a) The observed profile is converted into line records with position, integrated "
        "area and FWHM. (b) Match combines coverage of the strongest observed anchors with support of the candidate's "
        "strong reflections after bounded alignment. (c) Accepted phase profiles are jointly fitted, exposing both "
        "uncovered lines and deficits at shared reflections. (d) Gain searches direct, overlapping and hidden-phase "
        "evidence, rewards supported additional signal, and penalizes duplicate or unsupported strong reflections."
    ),
    (
        "Before any phase is accepted, the list is ordered mainly by Match. The purpose of this first pass is to recover "
        "the dominant phase without requiring every experimental line to belong to it. In the Ca2Al2SiO7 example, several "
        "COD entries have similar chemistry and space-group information. The line-integral score places COD#100048 "
        "Gehlenite at Match = 83% and Gain = 21%, followed closely by COD#4124699 Al2Ca2O7Si at Match = 81% and Gain = "
        "22%. Their I/Ic values are both close to 2.0, so the ranking alone does not resolve the assignment. The user must "
        "still check the source record, coverage of the leading reflections and strong calculated lines that are absent "
        "from the experiment; one coincident maximum is not enough."
    ): (
        "Before any phase is accepted, the list is ordered by Match and the Gain column is intentionally left empty. The "
        "purpose of this first pass is to recover the dominant phase without requiring every experimental line to belong "
        "to it. In the Ca2Al2SiO7 example, several COD entries have similar chemistry and space-group information. In the "
        "version 1.1.4 baseline run, COD#100048 Gehlenite reached Match = 83%, followed closely by COD#4124699 "
        "Al2Ca2O7Si at Match = 81%. Their I/Ic values are both close to 2.0, so the ranking alone does not resolve the "
        "assignment. The user must still check the source record, coverage of the leading reflections and strong "
        "calculated lines that are absent from the experiment; one coincident maximum is not enough."
    ),
    (
        "Dominant-phase identification examples. (a) Calcium aluminosilicate pattern with COD Ca2Al2SiO7-related "
        "candidates ranked by line-integral Match: COD#100048 Gehlenite gives Match = 83%, Gain = 21% and I/Ic = 2.05, "
        "while COD#4124699 Al2Ca2O7Si gives Match = 81%, Gain = 22% and I/Ic = 1.99."
    ): (
        "Dominant-phase identification example from the version 1.1.4 baseline run. Calcium aluminosilicate candidates "
        "are ordered by Match before any phase is accepted: COD#100048 Gehlenite gives Match = 83% and I/Ic = 2.05, "
        "while COD#4124699 Al2Ca2O7Si gives Match = 81% and I/Ic = 1.99. In version 1.2.0, Gain is intentionally shown "
        "only after a phase has been accepted."
    ),
    (
        "Current 50-case prototype validation matrix. The values describe the present implementation and candidate pool; "
        "they are not final universal accuracy claims."
    ): (
        "Prototype validation matrix for the version 1.1.4 scoring baseline. These values do not validate the revised "
        "version 1.2.0 Gain implementation and are not universal accuracy claims."
    ),
    "Current 50-case prototype validation": "Version 1.1.4 prototype validation baseline",
    (
        "Gain currently uses conservative thresholds: residual evidence must exceed a robust three-sigma background "
        "estimate and must include several independent line matches. Under these settings, none of the 30 closed "
        "single-phase RRUFF tests produced a false residual Gain above 5%. Relaxing the thresholds may recover more "
        "overlapped or textured minor phases, but it also increases the risk of treating preferred orientation, "
        "profile-width mismatch or background residuals as a new phase. The difference curve and unexplained-line markers "
        "are therefore retained as the evidence for deciding whether a broader search is justified."
    ): (
        "The version 1.1.4 Gain baseline used conservative thresholds: residual evidence had to exceed a robust "
        "three-sigma background estimate and include several independent line matches. Under those settings, none of the "
        "30 closed single-phase RRUFF tests produced a false residual Gain above 5%. Version 1.2.0 adds joint profile "
        "refitting and direct, overlap and hidden evidence stages; its thresholds and false-addition rate must therefore "
        "be re-evaluated rather than inferred from the earlier result."
    ),
    (
        "Table 2 summarizes the current validation run. The numbers should be read as the observed behaviour of version "
        "1.1.4 under the present scoring settings and candidate pool. Future work may improve Gain by adding an explicit "
        "relaxed residual-search mode, better handling of preferred orientation and broader calibration against "
        "independently measured mixtures."
    ): (
        "Table 2 summarizes the version 1.1.4 baseline run. These numbers remain useful for comparison, but they are not "
        "reported as performance of the revised version 1.2.0 scoring logic. Before submission, the same archived "
        "patterns, candidate pool and database snapshot should be rerun with version 1.2.0, followed by independently "
        "interpreted mixtures that challenge peak overlap, preferred orientation and minor-phase recovery."
    ),
    (
        "XRD Phase Finder is a cross-platform, open-source workspace for interpreting which crystalline phases are "
        "present in a powder mixture. It supports both an automatic search and a semi-automatic workflow in which the "
        "user reviews and accepts phases between iterations. COD access by itself is not new. The contribution of the "
        "program is the combination of several open structure sources, direct CIF input and an iterative search whose "
        "evidence remains visible. Match compares calculated sticks with the complete set of experimental line records. "
        "After a phase is accepted, Gain ranks the evidence left in the residual line set. The selected-phase fit also "
        "provides a clearly labelled semi-quantitative estimate from normalized non-negative profile contributions. "
        "Supporting and conflicting lines, processing choices and the current phase model remain part of the saved project."
    ): (
        "XRD Phase Finder is a cross-platform, open-source workspace for interpreting which crystalline phases are "
        "present in a powder mixture. It supports both an automatic search and a semi-automatic workflow in which the "
        "user reviews and accepts phases between iterations. COD access by itself is not new. The contribution of the "
        "program is the combination of several open structure sources, direct CIF input and an iterative search whose "
        "evidence remains visible. Match tests global two-way agreement between observed anchors and candidate "
        "reflections. After a phase is accepted, staged Gain asks what another candidate adds through uncovered, "
        "under-fitted overlapping or hidden-phase evidence. The selected-phase fit also provides a clearly labelled "
        "semi-quantitative estimate from normalized non-negative profile contributions. Supporting and conflicting "
        "lines, processing choices and the current phase model remain part of the saved project."
    ),
    (
        "In the 50 cases tested here, dominant-phase retrieval was stable for the RRUFF single-phase set and the "
        "conservative Gain settings avoided false additions above 5%. Weak and strongly overlapped impurities were less "
        "consistent and still required inspection of the difference curve. The next validation step is a larger public "
        "benchmark with corundum-normalized impurity levels, conservative and relaxed Gain settings, and evaluation by "
        "users who were not involved in development."
    ): (
        "The version 1.1.4 baseline recovered the dominant phase reliably in the tested RRUFF single-phase set and its "
        "conservative residual thresholds avoided false additions above 5%. Because version 1.2.0 changes both Match and "
        "Gain, those figures must not be transferred to the current implementation without a repeated benchmark. The "
        "next validation step is a public version 1.2.0 dataset with corundum-normalized impurity levels, explicit overlap "
        "cases and evaluation by users who were not involved in development."
    ),
}


def _replace_figure(document: Document) -> int:
    replacements = 0
    for shape in document.inline_shapes:
        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        relationship_id = blip.get(qn("r:embed"))
        relationship = document.part.rels[relationship_id]
        if relationship.target_ref != "media/image3.png":
            continue
        relationship.target_part._blob = MATCH_GAIN_FIGURE.read_bytes()
        with Image.open(MATCH_GAIN_FIGURE) as image:
            width_px, height_px = image.size
        shape.height = round(shape.width * height_px / width_px)
        replacements += 1
    return replacements


def update_manuscript(source: Path, output: Path) -> None:
    document = Document(source)
    replaced = 0
    for paragraph in document.paragraphs:
        text = paragraph.text
        replacement = REPLACEMENTS.get(text)
        if replacement is None and text.startswith(
            "XRD Phase Finder version 1.1.4 is implemented in Python 3.11–3.12."
        ):
            replacement = text.replace(
                "XRD Phase Finder version 1.1.4 is implemented in Python 3.11–3.12.",
                "XRD Phase Finder version 1.2.0 is implemented in Python 3.11–3.12.",
                1,
            )
        if replacement is None:
            continue
        paragraph.text = replacement
        replaced += 1

    if replaced != len(REPLACEMENTS):
        missing = [source_text for source_text in REPLACEMENTS if not any(p.text == REPLACEMENTS[source_text] for p in document.paragraphs)]
        raise RuntimeError(f"Replaced {replaced}/{len(REPLACEMENTS)} paragraphs; missing {len(missing)}")
    if _replace_figure(document) != 1:
        raise RuntimeError("Match/Gain figure was not replaced exactly once")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Updated {replaced} paragraphs and the Match/Gain figure: {output}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    update_manuscript(args.source, args.output)


if __name__ == "__main__":
    main()
