from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.table import Table


PARAGRAPH_REPLACEMENTS = {
    "XRD Phase Finder version 1.2.0 is implemented in Python 3.11–3.12.": (
        "XRD Phase Finder version 1.2.2 is implemented in Python 3.11–3.12."
    ),
    "Validation matrix for the version 1.2.0 Match and Gain implementation.": (
        "Validation matrix for the version 1.2.2 Match and Gain implementation."
    ),
    "Version 1.2.0 closed-world validation": "Version 1.2.2 closed-world validation",
    "The version 1.2.0 implementation was evaluated on 50 synthetic profiles": (
        "The version 1.2.2 implementation was evaluated on 50 synthetic profiles"
    ),
}


FULL_PARAGRAPH_REPLACEMENTS = {
    (
        "The useful comparison is therefore the complete workflow rather than a "
        "single feature. XRD Phase Finder combines open-source distribution, "
        "candidate generation from several experimental and computational "
        "repositories (Gražulis et al., 2009; Gražulis et al., 2012; Jain et al., "
        "2013; Curtarolo et al., 2012; Saal et al., 2013; Lafuente et al., 2015), "
        "direct CIF input, visible peak-level and residual evidence, persistent "
        "projects, and Gain ranking after accepted phases have been fitted. The "
        "evaluation below tests the added value of Gain against Match-only ranking. "
        "A broader comparison with open and commercial search–match workflows "
        "(Altomare et al., 2015; Lv et al., 2024; Lutterotti et al., 2019; Arcelus "
        "et al., 2024; Cao, 2026; Putz & Brandenburg, 2024; Degen et al., 2014; "
        "International Centre for Diffraction Data, 2026; Bruker, 2026; STOE & Cie "
        "GmbH, 2026) will require matched database snapshots and, for licensed "
        "software, reproducible access to the same reference data."
    ): (
        "The useful comparison is therefore the complete workflow rather than a "
        "single feature. XRD Phase Finder combines open-source distribution, "
        "candidate generation from several experimental and computational "
        "repositories (Gražulis et al., 2009; Gražulis et al., 2012; Jain et al., "
        "2013; Curtarolo et al., 2012; Saal et al., 2013; Lafuente et al., 2015), "
        "direct CIF input, visible peak-level and residual evidence, persistent "
        "projects, and Gain ranking after accepted phases have been fitted. The "
        "evaluation below compares Gain with internal Match-only and residual-ranking "
        "baselines. Published open and commercial programs (Altomare et al., 2015; "
        "Lv et al., 2024; Lutterotti et al., 2019; Arcelus et al., 2024; Cao, 2026; "
        "Putz & Brandenburg, 2024; Degen et al., 2014; International Centre for "
        "Diffraction Data, 2026; Bruker, 2026; STOE & Cie GmbH, 2026) are used to "
        "position the workflow, not as numerical benchmarks. A direct comparison "
        "would require matched database snapshots and reproducible access to the "
        "same licensed reference data."
    ),
    (
        "Conditional Gain was evaluated over 40 expected phase-addition steps from "
        "the binary and ternary profiles. After the already accepted families were "
        "fitted, the next expected family ranked first by Gain in 29/40 steps and "
        "appeared in the top five in 36/40; it received positive Gain in 37/40. "
        "These results support the use of Gain as a residual shortlist, while the "
        "remaining failures show that extensive overlap and weak components still "
        "require manual inspection."
    ): (
        "Conditional Gain was evaluated over 40 expected phase-addition steps from "
        "the binary and ternary profiles. After the accepted families were fitted, "
        "the next expected family ranked first by Gain in 23/40 steps and appeared "
        "in the top five in 31/40; it received positive Gain in 31/40. The version "
        "1.2.2 evidence gates are deliberately conservative: they reduce unsupported "
        "additions, but they also reject some weak or strongly overlapping true "
        "components. Gain should therefore be read as a residual shortlist rather "
        "than as an exhaustive detector."
    ),
    (
        "The same benchmark included 20 closed single-phase profiles for which any "
        "additional family was false. One profile produced a false Gain of at least "
        "5%, while the median maximum false Gain was 0%. Thus, the multiple-line, "
        "missing-line and profile-improvement gates substantially reduced unsupported "
        "additions but did not make them impossible."
    ): (
        "The same benchmark included 20 closed single-phase profiles for which any "
        "additional family was false. None produced a false Gain of at least 5%, and "
        "the median maximum false Gain was 0.09%. In this controlled candidate pool, "
        "the multiple-line, missing-line and profile-improvement gates prevented "
        "threshold-crossing false additions. This result describes the closed test "
        "set and should not be interpreted as a universal false-positive rate."
    ),
    (
        "A separate noise-significance stress test contained 40 binary profiles and "
        "ten single-phase negative controls. The signal-to-noise ratio was defined as "
        "the height of the second-strongest true minor-phase line divided by the known "
        "pointwise Gaussian-plus-counting noise standard deviation. The expected "
        "family appeared in the Gain top five in 2/13 cases below SNR 2, 2/11 cases "
        "at SNR 2-3, 3/7 cases at SNR 3-5 and 8/9 cases at SNR >= 5. Six of the nine "
        "SNR >= 5 cases ranked first. Two of the ten negative controls nevertheless "
        "produced false Gain above 5% (maximum 8.35%)."
    ): (
        "A separate noise-significance stress test contained 40 binary profiles and "
        "ten single-phase negative controls. The signal-to-noise ratio was defined as "
        "the height of the second-strongest true minor-phase line divided by the known "
        "pointwise Gaussian-plus-counting noise standard deviation. The expected "
        "family appeared in the Gain top five in 4/13 cases below SNR 2, 1/11 cases "
        "at SNR 2–3, 2/7 cases at SNR 3–5 and 3/9 cases at SNR ≥5. Two of the nine "
        "SNR ≥5 cases ranked first. None of the ten negative controls produced Gain "
        "above 5%; the maximum false Gain was 1.43%."
    ),
    (
        "Table 2 summarizes both tests. The sharp dependence on SNR is scientifically "
        "important: a three-sigma line threshold controls which reflections are "
        "testable, but it does not guarantee correct phase identification. Signals "
        "near the noise floor are retained as tentative candidates rather than "
        "converted into automatic positive calls. The reported percentages describe "
        "ranking positions and fitted profile contributions, not mass detection "
        "limits. Validation on independently interpreted laboratory mixtures remains "
        "necessary."
    ): (
        "Table 2 summarizes both tests. Recovery did not increase monotonically with "
        "this single-line SNR measure. The three-sigma threshold determines whether an "
        "individual reflection is testable, whereas Gain also requires several "
        "consistent lines, penalizes unsupported calculated reflections and accounts "
        "for overlap with accepted phases. SNR alone therefore does not determine the "
        "rank of a minor phase. The reported percentages describe ranking positions "
        "and fitted profile contributions, not mass detection limits. Validation on "
        "independently interpreted laboratory mixtures remains necessary."
    ),
    (
        "In the version 1.2.0 closed-world benchmark, Match placed the dominant formula "
        "family first in 48 of 50 profiles and within the top five in all 50. "
        "Conditional Gain placed the next expected family in the top five in 36 of 40 "
        "phase-addition steps. The noise-significance test showed that useful "
        "minor-phase ranking emerged mainly when more than one characteristic line "
        "rose clearly above the noise. These results support the intended use of the "
        "program as an inspectable automatic or semi-automatic phase-hypothesis tool. "
        "They do not establish universal accuracy, analytical detection limits or "
        "quantitative mass fractions. Independent laboratory evaluation and an "
        "archived database snapshot remain the next validation steps."
    ): (
        "In the version 1.2.2 closed-world benchmark, Match placed the dominant formula "
        "family first in 48 of 50 profiles and within the top five in all 50. "
        "Conditional Gain placed the next expected family in the top five in 31 of 40 "
        "phase-addition steps. No Gain of at least 5% was produced for the 20 closed "
        "single-phase profiles or the ten negative controls. These results support "
        "the intended use of the program as an inspectable automatic or semi-automatic "
        "phase-hypothesis tool, with a conservative residual search that favors "
        "specificity over exhaustive recovery of weak components. They do not "
        "establish universal accuracy, analytical detection limits or quantitative "
        "mass fractions. Independent laboratory evaluation and an archived database "
        "snapshot remain the next validation steps."
    ),
}


TABLE_UPDATES = {
    "Conditional Gain": (
        "40 expected phase-addition steps from the binary and ternary profiles",
        "Expected family: top-1 = 23/40; top-5 = 31/40; positive Gain = 31/40",
        "The conservative residual ranking promoted most true additions but rejected "
        "some weak or strongly overlapping components.",
    ),
    "False additions": (
        "20 closed single-phase profiles",
        "False Gain ≥5% = 0/20; median maximum false Gain = 0.09%",
        "No threshold-crossing false addition occurred in this closed test set.",
    ),
    "Noise-significance stress test": (
        "40 binary profiles and 10 negative controls",
        "Gain top-5 by secondary-line SNR: 4/13 (<2), 1/11 (2–3), 2/7 "
        "(3–5), 3/9 (≥5); false Gain ≥5%: 0/10 controls; maximum false "
        "Gain = 1.43%",
        "SNR alone did not determine rank because Gain also requires consistent "
        "multi-line evidence and penalizes unsupported reflections.",
    ),
}


def replace_paragraph_text(paragraph, old: str, new: str) -> bool:
    text = paragraph.text
    if old not in text:
        return False

    replacement = text.replace(old, new)
    if not paragraph.runs:
        paragraph.add_run(replacement)
        return True

    paragraph.runs[0].text = replacement
    for run in paragraph.runs[1:]:
        run.text = ""
    return True


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def add_ablation_section(document: Document) -> None:
    heading = document.add_paragraph(
        "Ablation of residual ranking",
        style="IUCr heading 2",
    )
    introduction = document.add_paragraph(
        "The four ranking modes were compared on the same 40 conditional "
        "phase-addition steps. Original Match retained the full-pattern ranking "
        "after accepted families were excluded. Residual Match repeated Match on "
        "peaks detected in the positive residual. Production Gain added the "
        "multi-line evidence gate, profile-improvement test and unsupported-line "
        "penalties. Raw Gain measured residual profile coverage without these "
        "safeguards.",
        style="IUCr body text",
    )
    caption = document.add_paragraph(
        "Ablation of conditional ranking in version 1.2.2. Top-k and mean "
        "reciprocal rank (MRR) refer to the next expected formula family. "
        "False-addition counts use the 20 closed single-phase controls; Match "
        "scores do not have a directly comparable automatic acceptance threshold.",
        style="IUCr table caption",
    )

    table_xml = deepcopy(document.tables[1]._tbl)
    while len(table_xml.tr_lst) > 5:
        table_xml.remove(table_xml.tr_lst[-1])
    table = Table(table_xml, document)
    rows = (
        (
            "Mode",
            "Ranking result",
            "Acceptance behaviour",
            "Interpretation",
        ),
        (
            "Original Match",
            "Top-1 = 15/40; top-5 = 31/40; MRR = 0.541",
            "No comparable Gain threshold",
            "Global ranking provides the baseline after accepted families are excluded.",
        ),
        (
            "Residual Match",
            "Top-1 = 30/40; top-5 = 37/40; MRR = 0.809",
            "No comparable Gain threshold",
            "Residual peak ranking gave the strongest shortlist in this closed test.",
        ),
        (
            "Production Gain",
            "Top-1 = 23/40; top-5 = 31/40; MRR = 0.653",
            "Positive = 31/40; false ≥5% = 0/20",
            "Evidence gates traded sensitivity for specificity.",
        ),
        (
            "Raw Gain",
            "Top-1 = 29/40; top-5 = 35/40; MRR = 0.794",
            "Positive = 38/40; false ≥5% = 14/20",
            "Removing gates recovered more true additions but admitted many false ones.",
        ),
    )
    for row, values in zip(table.rows, rows, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            set_cell_text(cell, value)

    interpretation = document.add_paragraph(
        "Residual Match produced the best top-five recovery, whereas production "
        "Gain was more selective. Removing the Gain gates restored much of the "
        "ranking sensitivity but produced a score of at least 5% in 14 of 20 "
        "single-phase controls. The ablation therefore does not show that Gain "
        "dominates every residual-ranking baseline. It shows the purpose of the "
        "current Gain design: to attach an explicit penalty to unsupported "
        "additions while keeping the evidence visible for manual review.",
        style="IUCr body text",
    )

    anchor = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("Table 2 summarizes both tests.")
    )._p
    for element in (
        heading._p,
        introduction._p,
        caption._p,
        table_xml,
        interpretation._p,
    ):
        anchor.addnext(element)
        anchor = element


def update_document(source: Path, target: Path) -> None:
    document = Document(source)
    changes = 0

    for paragraph in document.paragraphs:
        original = paragraph.text
        if original in FULL_PARAGRAPH_REPLACEMENTS:
            replace_paragraph_text(
                paragraph, original, FULL_PARAGRAPH_REPLACEMENTS[original]
            )
            changes += 1
            continue

        for old, new in PARAGRAPH_REPLACEMENTS.items():
            if replace_paragraph_text(paragraph, old, new):
                changes += 1

    for table in document.tables:
        for row in table.rows[1:]:
            label = row.cells[0].text.strip()
            if label not in TABLE_UPDATES:
                continue
            material, result, interpretation = TABLE_UPDATES[label]
            set_cell_text(row.cells[1], material)
            set_cell_text(row.cells[2], result)
            set_cell_text(row.cells[3], interpretation)
            changes += 1

    add_ablation_section(document)

    if changes != 13:
        raise RuntimeError(f"Expected 13 document updates, applied {changes}")

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("target", type=Path)
    args = parser.parse_args()
    update_document(args.source, args.target)


if __name__ == "__main__":
    main()
