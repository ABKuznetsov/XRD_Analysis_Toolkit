from __future__ import annotations

import argparse
import re
import shutil
import tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


CITATIONS = {
    1: "Dinnebier & Billinge, 2008",
    2: "Gražulis et al., 2009",
    3: "Gražulis et al., 2012",
    4: "Altomare et al., 2015",
    5: "Lv et al., 2024",
    6: "Lutterotti et al., 2019",
    7: "Arcelus et al., 2024",
    8: "Cao, 2026",
    9: "Putz & Brandenburg, 2024",
    10: "Degen et al., 2014",
    11: "International Centre for Diffraction Data, 2026",
    12: "Bruker, 2026",
    13: "STOE & Cie GmbH, 2026",
    14: "Jain et al., 2013",
    15: "Curtarolo et al., 2012",
    16: "Saal et al., 2013",
    17: "Lafuente et al., 2015",
    18: "Toby & Von Dreele, 2013",
    19: "Rodríguez-Carvajal, 1993",
    20: "Coelho, 2018",
    21: "Momma & Izumi, 2011",
    22: "Macrae et al., 2008",
    23: "Wojdyr, 2022",
    24: "Ong et al., 2013",
    25: "Togo et al., 2024",
    26: "Larsen et al., 2017",
    27: "Harris et al., 2020",
    28: "Virtanen et al., 2020",
    29: "McKinney, 2010",
    30: "Hunter, 2007",
    31: "The Qt Company, 2026",
    34: "Savitzky & Golay, 1964",
    35: "Eilers & Boelens, 2005",
    36: "Baek et al., 2015",
    37: "Morháč et al., 1997",
    38: "Erb, 2022",
    39: "Schwarz, 1978",
    40: "Cromer & Mann, 1968",
    41: "Gilmore et al., 2019",
    42: "Thompson et al., 1987",
    44: "Lawson & Hanson, 1995",
}

DROP_REFERENCES = {32, 33, 43}

SPGLIB_REFERENCE = (
    "Togo, A., Shinohara, K. & Tanaka, I. (2024). "
    "Spglib: a software library for crystal symmetry search. "
    "Sci. Technol. Adv. Mater. Methods 4, 2384822. "
    "https://doi.org/10.1080/27660400.2024.2384822."
)

ABSTRACT = (
    "XRD Phase Finder is an open-source desktop application for automatic and "
    "semi-automatic interpretation of the crystalline phases present in powder "
    "mixtures. Automatic search generates and ranks candidates, whereas the "
    "semi-automatic workflow allows the user to set elemental constraints, inspect "
    "the diffraction evidence and accept phases between residual-search iterations. "
    "The program searches user-accessible records from the Crystallography Open "
    "Database, Materials Project, AFLOW, OQMD and RRUFF, and it can evaluate an "
    "individual CIF file without adding it to a database. Candidate patterns are "
    "calculated locally and displayed with the experimental profile, reflection "
    "markers, background components and unexplained lines. Match ranks initial "
    "candidates against the complete set of experimental line records. After one or "
    "more phases have been accepted, Gain re-ranks the remaining candidates using "
    "the residual line set and demotes entries that duplicate already explained "
    "reflections. Smoothing, physical-background estimation and an optional broad "
    "amorphous contribution remain inspectable and reversible. For accepted phases, "
    "normalized non-negative fitted profile contributions provide a semi-quantitative "
    "estimate. Projects retain imported patterns, processing settings, accepted "
    "phases and view state. Standalone Windows and macOS packages are distributed "
    "under the MIT licence. Match and Gain are ranking scores rather than phase "
    "probabilities, and the semi-quantitative values are not validated mass fractions."
)

SEMI_QUANT_METHOD = (
    "After phases have been accepted, their calculated profiles are fitted "
    "simultaneously to the active experimental signal by weighted non-negative least "
    "squares (Lawson & Hanson, 1995). The resulting non-negative coefficients are "
    "normalized across the selected phase set and displayed as Quant. (%). They are "
    "reported as semi-quantitative profile contributions. The current normalization "
    "does not correct for absorption, microabsorption, preferred orientation or "
    "phase-specific diffraction response. I/Ic is displayed separately and is not "
    "used in this normalization."
)

LIMITATION_SEMI_QUANT = (
    "The selected-phase table provides a semi-quantitative estimate by normalizing "
    "non-negative fitted profile contributions. The values are useful for comparing "
    "accepted phases within the current model, but they are not validated mass "
    "fractions because absorption, microabsorption, preferred orientation and "
    "phase-specific diffraction response are not fully corrected."
)

CONCLUSION_FIRST = (
    "XRD Phase Finder is a cross-platform, open-source workspace for interpreting "
    "which crystalline phases are present in a powder mixture. It supports both an "
    "automatic search and a semi-automatic workflow in which the user reviews and "
    "accepts phases between iterations. COD access by itself is not new. The "
    "contribution of the program is the combination of several open structure "
    "sources, direct CIF input and an iterative search whose evidence remains "
    "visible. Match compares calculated sticks with the complete set of experimental "
    "line records. After a phase is accepted, Gain ranks the evidence left in the "
    "residual line set. The selected-phase fit also provides a clearly labelled "
    "semi-quantitative estimate from normalized non-negative profile contributions. "
    "Supporting and conflicting lines, processing choices and the current phase "
    "model remain part of the saved project."
)

DATA_AVAILABILITY = (
    "The source code and release packages are available at "
    "https://github.com/ABKuznetsov/XRD_Analysis_Toolkit. The exact release used for "
    "the reported evaluation, benchmark patterns, ground-truth labels and analysis "
    "settings will be archived in a persistent repository before submission. Complete "
    "third-party structure collections are not bundled with the installer; users "
    "download and index the open sources required for their work or add individual "
    "CIF files directly. Reproduction of a benchmark must therefore report the exact "
    "database snapshot or archived candidate list used for the run."
)

AI_DISCLOSURE = (
    "Generative artificial-intelligence tools were used for language editing and "
    "schematic layout. The author checked all scientific statements, numerical "
    "results, source citations and final wording."
)


def _replace_package_parts(source: Path, template: Path, output: Path) -> None:
    replacements = {}
    with ZipFile(template) as template_zip:
        for name in ("word/styles.xml", "word/numbering.xml"):
            replacements[name] = template_zip.read(name)

    with ZipFile(source) as source_zip, ZipFile(output, "w", ZIP_DEFLATED) as target_zip:
        for item in source_zip.infolist():
            payload = replacements.get(item.filename, source_zip.read(item.filename))
            target_zip.writestr(item, payload)


def _remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def _clear_direct_run_formatting(paragraph) -> None:
    for run in paragraph.runs:
        r_pr = run._r.rPr
        if r_pr is not None:
            run._r.remove(r_pr)


def _expand_reference_token(token: str) -> list[int]:
    numbers: list[int] = []
    for part in re.split(r"[,;]", token):
        part = part.strip()
        range_match = re.fullmatch(r"(\d+)[–-](\d+)", part)
        if range_match:
            start, end = map(int, range_match.groups())
            numbers.extend(range(start, end + 1))
        elif part.isdigit():
            numbers.append(int(part))
    return numbers


def _convert_numeric_citations(text: str) -> str:
    pattern = re.compile(r"\[((?:\d+\s*(?:[,;–-]\s*)?)+)\]")

    def replace(match: re.Match[str]) -> str:
        numbers = _expand_reference_token(match.group(1))
        if not numbers or any(number not in CITATIONS for number in numbers):
            return match.group(0)
        labels = []
        for number in numbers:
            label = CITATIONS[number]
            if label not in labels:
                labels.append(label)
        return f"({'; '.join(labels)})"

    return pattern.sub(replace, text)


def _insert_before(paragraph, text: str, style: str):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    inserted = paragraph._parent.add_paragraph()
    inserted._p.getparent().remove(inserted._p)
    inserted._p = new_p
    inserted._element = new_p
    inserted._parent = paragraph._parent
    inserted.style = style
    inserted.add_run(text)
    return inserted


def _reference_sort_key(text: str) -> str:
    cleaned = re.sub(r"^[^\wÀ-ž]+", "", text)
    return cleaned.casefold()


def _set_cell_margins(cell, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = f"w:{edge}"
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_submission(source: Path, template: Path, output: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="jac_format_") as temp_dir:
        staged = Path(temp_dir) / "styled.docx"
        _replace_package_parts(source, template, staged)
        document = Document(staged)

        paragraphs = document.paragraphs
        by_text = {paragraph.text.strip(): paragraph for paragraph in paragraphs if paragraph.text.strip()}

        # Scientific wording aligned with the implemented calculation.
        for paragraph in paragraphs:
            text = paragraph.text.strip()
            if text.startswith("XRD Phase Finder is an open-source desktop application for interpreting"):
                paragraph.text = ABSTRACT
            elif text == (
                "Quantitative phase analysis is outside the scope of the current version. "
                "Values shown in the selected-phase table summarize fitted profile "
                "contributions within the program and must not be reported as validated "
                "mass fractions."
            ):
                paragraph.text = LIMITATION_SEMI_QUANT
            elif text.startswith(
                "XRD Phase Finder is a cross-platform, open-source workspace for interpreting"
            ):
                paragraph.text = CONCLUSION_FIRST
            elif text.startswith("The source code and release packages are available at"):
                paragraph.text = DATA_AVAILABILITY

        reproducibility = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.strip().endswith("Reproducibility and software availability")
        )
        _insert_before(reproducibility, "Semi-quantitative phase contributions", "IUCr heading 2")
        _insert_before(reproducibility, SEMI_QUANT_METHOD, "IUCr body text")

        # Convert the manuscript to author-year citations before the references are rebuilt.
        references_heading = next(
            paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "References"
        )
        for paragraph in document.paragraphs:
            if paragraph._p is references_heading._p:
                break
            converted = _convert_numeric_citations(paragraph.text)
            if converted != paragraph.text:
                paragraph.text = converted

        # Replace the obsolete spglib preprint and drop references no longer cited.
        reference_paragraphs = []
        after_references = False
        references: dict[int, str] = {}
        for paragraph in list(document.paragraphs):
            if paragraph._p is references_heading._p:
                after_references = True
                continue
            if not after_references:
                continue
            match = re.match(r"^\[(\d+)\]\s*(.*)$", paragraph.text.strip(), re.DOTALL)
            if not match:
                continue
            number = int(match.group(1))
            text = SPGLIB_REFERENCE if number == 25 else match.group(2).strip()
            if number not in DROP_REFERENCES:
                references[number] = text
            reference_paragraphs.append(paragraph)

        for paragraph in reference_paragraphs:
            _remove_paragraph(paragraph)

        for text in sorted(references.values(), key=_reference_sort_key):
            paragraph = document.add_paragraph(text, style="IUCr references")
            _clear_direct_run_formatting(paragraph)

        # Back matter follows the dedicated IUCr styles where available.
        abstract_heading = next(
            (p for p in document.paragraphs if p.text.strip() == "Abstract"), None
        )
        if abstract_heading is not None:
            _remove_paragraph(abstract_heading)

        funding_heading = next(
            p for p in document.paragraphs if p.text.strip() == "Acknowledgements"
        )
        funding_heading.text = "Funding information"
        funding_heading.style = "IUCr heading 4"
        funding_body = funding_heading._p.getnext()
        funding_paragraph = next(p for p in document.paragraphs if p._p is funding_body)
        funding_paragraph.style = "IUCr body text"
        ai_paragraph = funding_paragraph.insert_paragraph_before(AI_DISCLOSURE)
        funding_paragraph._p.addnext(ai_paragraph._p)
        ai_paragraph.style = "IUCr acknowledgements"

        conflict_heading = next(
            p for p in document.paragraphs if p.text.strip() == "Conflict of interest"
        )
        conflict_body_node = conflict_heading._p.getnext()
        conflict_body = next(p for p in document.paragraphs if p._p is conflict_body_node)
        conflict_body.style = "IUCr coi"
        _remove_paragraph(conflict_heading)

        data_heading = next(
            p for p in document.paragraphs if p.text.strip() == "Data and code availability"
        )
        data_body_node = data_heading._p.getnext()
        data_body = next(p for p in document.paragraphs if p._p is data_body_node)
        data_body.style = "IUCr da"
        _remove_paragraph(data_heading)

        # Apply the official IUCr paragraph styles and remove manual labels that
        # are generated by those styles.
        for index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                if paragraph._p.xpath(".//w:drawing"):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                continue

            if index == 0:
                paragraph.style = "IUCr article title"
            elif text == "A. B. Kuznetsov":
                paragraph.style = "IUCr authors"
            elif text.startswith("Sobolev Institute") or text.startswith("Correspondence:"):
                paragraph.style = "IUCr affiliations"
            elif text.startswith("Synopsis. "):
                paragraph.text = text.removeprefix("Synopsis. ")
                paragraph.style = "IUCr synopsis"
            elif text == ABSTRACT:
                paragraph.style = "IUCr abstract"
            elif text.startswith("Keywords:"):
                paragraph.text = text.removeprefix("Keywords:").strip()
                paragraph.style = "IUCr keywords"
            elif re.match(r"^\d+\.\s+", text):
                paragraph.text = re.sub(r"^\d+\.\s+", "", text)
                paragraph.style = "IUCr heading 1"
            elif re.match(r"^\d+\.\d+\.\s+", text):
                paragraph.text = re.sub(r"^\d+\.\d+\.\s+", "", text)
                paragraph.style = "IUCr heading 2"
            elif re.match(r"^Figure\s+\d+\.\s*", text):
                paragraph.text = re.sub(r"^Figure\s+\d+\.\s*", "", text)
                paragraph.style = "IUCr figure caption"
            elif re.match(r"^Table\s+\d+\.\s*", text):
                paragraph.text = re.sub(r"^Table\s+\d+\.\s*", "", text)
                paragraph.style = "IUCr table caption"
            elif text == "References":
                paragraph.style = "IUCr heading 4"
            elif text in {"Funding information", "Author contributions"}:
                paragraph.style = "IUCr heading 4"
            elif paragraph.style.name not in {
                "IUCr synopsis",
                "IUCr abstract",
                "IUCr keywords",
                "IUCr acknowledgements",
                "IUCr coi",
                "IUCr da",
                "IUCr references",
                "IUCr heading 1",
                "IUCr heading 2",
                "IUCr heading 4",
                "IUCr article title",
                "IUCr authors",
                "IUCr affiliations",
                "IUCr figure caption",
                "IUCr table caption",
            }:
                paragraph.style = "IUCr body text"

            _clear_direct_run_formatting(paragraph)

        # Equations are short display lines, not ordinary prose.
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text.startswith("M = 100") or text.startswith("G = F("):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if paragraph.style.name in {
                "IUCr heading 1",
                "IUCr heading 2",
                "IUCr heading 4",
                "IUCr figure caption",
                "IUCr table caption",
            }:
                paragraph.paragraph_format.keep_with_next = True

        for table in document.tables:
            table.autofit = True
            for row_index, row in enumerate(table.rows):
                for cell in row.cells:
                    _set_cell_margins(cell)
                    for paragraph in cell.paragraphs:
                        paragraph.style = "IUCr table text"
                        _clear_direct_run_formatting(paragraph)
                        if row_index == 0:
                            for run in paragraph.runs:
                                run.bold = True

        # Match the generic IUCr template page geometry.
        template_document = Document(template)
        template_section = template_document.sections[0]
        for section in document.sections:
            section.page_width = template_section.page_width
            section.page_height = template_section.page_height
            section.top_margin = template_section.top_margin
            section.bottom_margin = template_section.bottom_margin
            section.left_margin = template_section.left_margin
            section.right_margin = template_section.right_margin
            section.start_type = WD_SECTION.NEW_PAGE if section is not document.sections[0] else section.start_type
            for footer in (section.footer, section.first_page_footer, section.even_page_footer):
                for paragraph in footer.paragraphs:
                    paragraph.text = "XRD Phase Finder"

        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Format the XRD Phase Finder manuscript for IUCr/JAC submission.")
    parser.add_argument("source", type=Path)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    format_submission(args.source, args.template, args.output)


if __name__ == "__main__":
    main()
